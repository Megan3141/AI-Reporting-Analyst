from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook


ROOT_FOLDERS = [
    Path("path/to/recovery_files"),
]

OUTPUT_FILE = Path("path/to/recovery_extract.xlsx")

SEARCH_TERMS = [
    "region_a",
    "region_b",
]

VALID_CATEGORIES = [
    "Category_A",
    "Category_B",
    "Category_C",
    "Category_D",
]


def extract_period(file_path):
    for part in file_path.parts:
        if part.upper().startswith("P") and part[1:].isdigit():
            return part.upper()

    return None


def extract_week(file_path):
    match = re.search(
        r"(?:WK|WEEK|WEEKLY)\s*0?(\d{1,2})",
        file_path.name,
        re.IGNORECASE,
    )

    if match:
        return int(match.group(1))

    return None


def find_info_sheet(workbook):
    preferred_sheets = [
        "Info - Split",
        "Info",
        "Info - Total",
    ]

    for sheet_name in preferred_sheets:
        if sheet_name in workbook.sheetnames:
            return sheet_name

    return None


def is_recovery_file(file_path):
    full_path = str(file_path).lower()
    file_name = file_path.name.lower()

    if file_path.suffix.lower() not in [".xlsx", ".xlsm"]:
        return False

    if file_name.startswith("~$"):
        return False

    if "std" in full_path:
        return False

    if "archive" in full_path:
        return False

    if not any(term in full_path for term in SEARCH_TERMS):
        return False

    return "wk" in file_name or "week" in file_name or "weekly" in file_name


def extract_file(file_path):
    rows = []
    log_entry = {
        "Period": extract_period(file_path),
        "Week": extract_week(file_path),
        "File": file_path.name,
        "Sheet Name": "",
        "Status": "",
        "Reason": "",
        "Rows Extracted": 0,
    }

    try:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        sheet_name = find_info_sheet(workbook)

        if sheet_name is None:
            log_entry["Status"] = "Skipped"
            log_entry["Reason"] = "No valid info sheet found"
            workbook.close()
            return rows, log_entry

        worksheet = workbook[sheet_name]
        log_entry["Sheet Name"] = sheet_name

        for row in range(1, worksheet.max_row + 1):
            category = worksheet[f"B{row}"].value
            dimension = worksheet[f"C{row}"].value
            level_1 = worksheet[f"D{row}"].value
            metric_1 = worksheet[f"F{row}"].value
            metric_2 = worksheet[f"AN{row}"].value
            metric_3 = worksheet[f"AO{row}"].value

            if category is not None:
                category = str(category).strip()

            if dimension is not None:
                dimension = str(dimension).strip()

            if not category:
                continue

            if not dimension:
                continue

            if "total" in category.lower():
                continue

            if category not in VALID_CATEGORIES:
                continue

            if metric_1 in [0, 0.0, None]:
                continue

            rows.append({
                "Period": log_entry["Period"],
                "Week": log_entry["Week"],
                "Source File": file_path.name,
                "Sheet Name": sheet_name,
                "Category": category,
                "Dimension": dimension,
                "Level 1": level_1,
                "Metric 1": metric_1,
                "Metric 2": metric_2,
                "Metric 3": metric_3,
            })

        workbook.close()

        log_entry["Status"] = "Success"
        log_entry["Rows Extracted"] = len(rows)

        return rows, log_entry

    except Exception as error:
        log_entry["Status"] = "Error"
        log_entry["Reason"] = str(error)

        return rows, log_entry


def extract_recovery_files():
    all_rows = []
    log_rows = []
    files = []

    for root_folder in ROOT_FOLDERS:
        for file in root_folder.rglob("*"):
            if file.is_file() and is_recovery_file(file):
                files.append(file)

    print("Recovery files found:", len(files))

    for file_path in files:
        print("Processing:", file_path.name)

        rows, log_entry = extract_file(file_path)

        all_rows.extend(rows)
        log_rows.append(log_entry)

    return pd.DataFrame(all_rows), pd.DataFrame(log_rows)


def export_results(dataframe, log_dataframe):
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for category in VALID_CATEGORIES:
            category_df = dataframe[dataframe["Category"] == category]

            if len(category_df) > 0:
                category_df.to_excel(
                    writer,
                    sheet_name=category[:31],
                    index=False,
                )

        log_dataframe.to_excel(
            writer,
            sheet_name="Extraction_Log",
            index=False,
        )


if __name__ == "__main__":
    df, log_df = extract_recovery_files()
    export_results(df, log_df)

    print()
    print("Recovery extract created:")
    print(OUTPUT_FILE)