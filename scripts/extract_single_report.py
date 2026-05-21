from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


SOURCE_FILE = Path("path/to/source_report.xlsx")
OUTPUT_FILE = Path("path/to/output_extract.xlsx")

VALID_CATEGORIES = [
    "Category_A",
    "Category_B",
    "Category_C",
    "Category_D",
]


def find_info_sheet(workbook):
    preferred_sheets = [
        "Info - Total",
        "Info - Split",
        "Info",
    ]

    for sheet_name in preferred_sheets:
        if sheet_name in workbook.sheetnames:
            return sheet_name

    return None


def extract_report(source_file):
    rows = []

    workbook = load_workbook(source_file, data_only=True, read_only=True)
    sheet_name = find_info_sheet(workbook)

    if sheet_name is None:
        workbook.close()
        raise ValueError("No valid info sheet found.")

    worksheet = workbook[sheet_name]

    for row in range(1, worksheet.max_row + 1):
        category = worksheet[f"B{row}"].value
        dimension = worksheet[f"C{row}"].value
        level_1 = worksheet[f"D{row}"].value
        metric_1 = worksheet[f"F{row}"].value
        metric_2 = worksheet[f"AN{row}"].value
        metric_3 = worksheet[f"AO{row}"].value

        if category is not None:
            category = str(category).strip()

        if dimension is not None:
            dimension = str(dimension).strip()

        if not category:
            continue

        if not dimension:
            continue

        if "total" in category.lower():
            continue

        if category not in VALID_CATEGORIES:
            continue

        if metric_1 in [0, 0.0, None]:
            continue

        rows.append({
            "Source File": source_file.name,
            "Sheet Name": sheet_name,
            "Category": category,
            "Dimension": dimension,
            "Level 1": level_1,
            "Metric 1": metric_1,
            "Metric 2": metric_2,
            "Metric 3": metric_3,
        })

    workbook.close()

    return pd.DataFrame(rows)


def export_by_category(dataframe, output_file):
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for category in VALID_CATEGORIES:
            category_df = dataframe[dataframe["Category"] == category]

            if len(category_df) > 0:
                category_df.to_excel(
                    writer,
                    sheet_name=category[:31],
                    index=False,
                )


if __name__ == "__main__":
    df = extract_report(SOURCE_FILE)
    export_by_category(df, OUTPUT_FILE)

    print("Single report extract created:")
    print(OUTPUT_FILE)